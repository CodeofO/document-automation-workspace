"""Shared request helpers and sample files for the API tests."""

import base64
import io
import time
import zipfile
from pathlib import Path
from typing import Any

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover - compatibility for older PyMuPDF installs
    import fitz


ONE_BY_ONE_PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII=")


SCHEMA_COUNTER = 0


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def assert_xlsx_response(response) -> None:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MIME
    assert response.content.startswith(b"PK")


def wait_for_export_job(client, job_id: str) -> dict[str, Any]:
    for _ in range(20):
        response = client.get(f"/api/export-jobs/{job_id}")
        assert response.status_code == 200, response.text
        payload = response.json()
        if payload["status"] in {"completed", "failed"}:
            return payload
        time.sleep(0.01)
    raise AssertionError(f"Export job did not finish: {job_id}")


def _read_env_example(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip()
    return values


def upload_png(client):
    response = client.post(
        "/api/documents",
        files={"file": ("invoice.png", ONE_BY_ONE_PNG, "image/png")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def create_schema(client, name: str | None = None):
    global SCHEMA_COUNTER
    if name is None:
        SCHEMA_COUNTER += 1
        schema_name = "invoice_basic" if SCHEMA_COUNTER == 1 else f"invoice_basic_{SCHEMA_COUNTER}"
    else:
        schema_name = name
    response = client.post(
        "/api/schemas",
        json={
            "name": schema_name,
            "display_name": schema_name.replace("_", " ").title(),
            "fields": [
                {
                    "key_name": "invoice_number",
                    "description": "Invoice number near the top of the document. Return null if missing.",
                    "output_format": "string",
                },
                {
                    "key_name": "total_amount",
                    "description": "Final total amount including tax.",
                    "output_format": "float",
                },
            ],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def create_document_classifier(client, name: str = "document_classifier"):
    response = client.post(
        "/api/document-classifiers",
        json={
            "name": name,
            "description": "문서를 사용자가 정의한 class 후보 중 하나로 분류합니다.",
            "allow_unknown": True,
            "classes": [
                {
                    "class_name": "contract",
                    "description": "계약 조건과 서명 또는 날인이 있는 문서",
                    "signals": ["계약", "서명", "날인"],
                },
                {
                    "class_name": "consent_form",
                    "description": "개인정보 또는 금융정보 조회 동의 여부가 있는 문서",
                    "signals": ["동의", "개인정보", "조회"],
                },
            ],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def create_required_field_checklist(client, name: str = "required_checklist"):
    response = client.post(
        "/api/required-field-checklists",
        json={
            "name": name,
            "description": "필수 항목의 존재 여부만 확인합니다.",
            "regions": [{"id": "signature_region", "name": "서명 영역", "page": 1, "x": 0.55, "y": 0.55, "width": 0.35, "height": 0.25}],
            "items": [
                {
                    "item_name": "성명",
                    "description": "성명이 문서에 존재하는지 확인합니다.",
                    "evidence_type": "text_or_handwriting",
                    "required": True,
                },
                {
                    "item_name": "서명",
                    "description": "서명 또는 날인이 존재하는지 확인합니다.",
                    "evidence_type": "signature_or_stamp",
                    "required": True,
                    "region_id": "signature_region",
                },
                {
                    "item_name": "체크박스",
                    "description": "필수 체크박스 표시가 존재하는지 확인합니다.",
                    "evidence_type": "checkbox",
                    "required": True,
                },
            ],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def workflow_definition(schema_id: str, classifier_id: str, checklist_id: str):
    return {
        "nodes": [
            {"id": "input", "position": {"x": 0, "y": 120}, "data": {"kind": "input", "label": "Input"}},
            {
                "id": "classifier",
                "position": {"x": 220, "y": 120},
                "data": {"kind": "classifier", "label": "Classifier", "config": {"classifier_id": classifier_id}},
            },
            {"id": "branch", "position": {"x": 440, "y": 120}, "data": {"kind": "branch", "label": "Branch"}},
            {
                "id": "kie_contract",
                "position": {"x": 680, "y": 60},
                "data": {"kind": "kie", "label": "Contract KIE", "config": {"schema_id": schema_id}},
            },
            {
                "id": "required_contract",
                "position": {"x": 900, "y": 60},
                "data": {"kind": "required-checker", "label": "Contract Required", "config": {"checklist_id": checklist_id}},
            },
            {"id": "merge", "position": {"x": 1120, "y": 120}, "data": {"kind": "merge", "label": "Merge"}},
            {"id": "export", "position": {"x": 1340, "y": 120}, "data": {"kind": "export", "label": "Export"}},
        ],
        "edges": [
            {"id": "input-classifier", "source": "input", "target": "classifier"},
            {"id": "classifier-branch", "source": "classifier", "target": "branch"},
            {"id": "branch-contract", "source": "branch", "target": "kie_contract", "sourceHandle": "class:contract"},
            {"id": "branch-consent-form", "source": "branch", "target": "merge", "sourceHandle": "class:consent_form"},
            {"id": "branch-unknown", "source": "branch", "target": "merge", "sourceHandle": "unknown"},
            {"id": "kie-required", "source": "kie_contract", "target": "required_contract"},
            {"id": "required-merge", "source": "required_contract", "target": "merge"},
            {"id": "merge-export", "source": "merge", "target": "export"},
        ],
    }


def make_pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page(width=240, height=120)
    page.insert_text((24, 60), "Invoice No. INV-2026-001")
    buffer = io.BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def make_pdf_with_image_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page(width=240, height=120)
    page.insert_text((24, 28), "Document with image")
    page.insert_image(fitz.Rect(24, 40, 80, 96), stream=ONE_BY_ONE_PNG)
    buffer = io.BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def make_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Quarterly Report</w:t></w:r></w:p>
    <w:p><w:r><w:t>Executive summary paragraph.</w:t></w:r></w:p>
    <w:tbl>
      <w:tr><w:tc><w:p><w:r><w:t>Metric</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Value</w:t></w:r></w:p></w:tc></w:tr>
      <w:tr><w:tc><w:p><w:r><w:t>Revenue</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr>
    </w:tbl>
  </w:body>
</w:document>"""
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def make_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Finance"
    sheet.append(["Metric", "Value"])
    sheet.append(["Revenue", 100])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def make_xlsx_formula_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Finance"
    sheet.append(["Metric", "Value"])
    sheet.append(["Revenue", 100])
    sheet.append(["Total", "=SUM(B2:B2)"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def make_pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Roadmap"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(6), Inches(1))
    textbox.text_frame.text = "Launch Raw Data Extractor"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_pptx_with_image_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Image slide"
    image_stream = io.BytesIO(ONE_BY_ONE_PNG)
    slide.shapes.add_picture(image_stream, Inches(1), Inches(1.4), width=Inches(1))
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
