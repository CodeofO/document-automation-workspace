"""Table extraction part 2: exports, workflow ``table`` node, AI column draft (recommendations)."""

import io
import zipfile
from typing import Any

from openpyxl import load_workbook
from PIL import Image

from app.application.table_exports import unique_sheet_name
from app.storage import read_storage_bytes
from tests.api.test_page_level_processing import upload_pdf
from tests.api.test_table_extraction import _create_template, _upload_image
from tests.support.api import create_document_classifier, create_schema, wait_for_export_job

LONG_NAME = "화학 성분/분석 결과 표 [측정 블록] 아주 긴 이름입니다"


def _injection_table() -> dict[str, Any]:
    return {
        "key": "injection_rows",
        "name": "주입 내역",
        "columns": [
            {"key": "품명"},
            {"key": "단중", "type": "number"},
            {"key": "실적", "type": "number"},
            {"key": "주입중량", "type": "number"},
            {"key": "레들", "type": "checkbox", "checkbox_options": ["A1", "A2", "A3"]},
        ],
        "rules": [{"kind": "row_formula", "expression": "주입중량 = 단중 * 실적"}],
    }


def _chem_table() -> dict[str, Any]:
    return {"key": "chem", "name": LONG_NAME, "columns": [{"key": "원소"}, {"key": "값", "type": "number", "keep_raw": False}]}


def _row(line: int, *cells: tuple[str | None, float]) -> dict[str, Any]:
    return {"source_line": line, "y_center": None, "cells": [{"raw": raw, "conf": conf} for raw, conf in cells]}


def _ledger_extract(table, image_inputs, *, has_region_crop):
    if table.key == "injection_rows":
        rows = [
            _row(1, ("MBS", 0.95), ("2", 0.95), ("3", 0.95), ("6", 0.95), ("A1:1, A3:1", 0.95)),
            _row(2, ("PP", 0.95), ("2", 0.7), ("3", 0.95), ("7", 0.95), ("A2", 0.95)),
        ]
    else:
        rows = [_row(1, ("C", 0.95), ("0.12", 0.95)), _row(2, ("Si", 0.95), ("1.5?", 0.9))]
    return {"headers_seen": [], "rows": rows, "notes": None}


def _sheet_rows(workbook, name: str) -> list[list[Any]]:
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


def _ledger_batch(fake_vlm, client) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    fake_vlm.set(extract_table=_ledger_extract)
    template = _create_template(client, name="샘플 주입일지", tables=[_injection_table(), _chem_table()])
    document = upload_pdf(client, filename="ledger.pdf", page_count=2)
    batch = client.post(
        "/api/table-extraction-batches/from-documents", json={"template_id": template["id"], "document_ids": [document["document_id"]]}
    )
    assert batch.status_code == 200, batch.text
    batch = client.get(f"/api/table-extraction-batches/{batch.json()['id']}").json()
    jobs = [client.get(f"/api/table-extraction-jobs/{item['job_id']}").json() for item in batch["items"]]
    assert [job["page_number"] for job in jobs] == [1, 2]
    # The reviewer fixes the formula violation on page 1 → corrected_tables take precedence in exports.
    patched = client.patch(
        f"/api/table-extraction-results/{jobs[0]['result_id']}",
        json={"edits": [{"table_key": "injection_rows", "row_index": 2, "column": "주입중량", "value": "6"}]},
    )
    assert patched.status_code == 200, patched.text
    return batch, jobs


# --- export ---------------------------------------------------------------------------------


def test_batch_export_xlsx_has_table_review_and_violation_sheets(mock_vlm, fake_vlm, client) -> None:
    batch, _jobs = _ledger_batch(fake_vlm, client)
    response = client.get(f"/api/table-extraction-batches/{batch['id']}/export?format=xlsx")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(io.BytesIO(response.content))
    chem_sheet = workbook.sheetnames[1]
    assert workbook.sheetnames == ["주입 내역", chem_sheet, "검토_칸", "검증_위반"]
    assert len(chem_sheet) <= 31 and "/" not in chem_sheet and "[" not in chem_sheet and chem_sheet.startswith("화학 성분_분석")

    assert _sheet_rows(workbook, "주입 내역") == [
        ["파일", "페이지", "행번호", "품명", "단중", "실적", "주입중량", "레들"],
        ["ledger.pdf", 1, 1, "MBS", 2, 3, 6, "A1;A3"],
        ["ledger.pdf", 1, 2, "PP", 2, 3, 6, "A2"],  # corrected value
        ["ledger.pdf", 2, 1, "MBS", 2, 3, 6, "A1;A3"],
        ["ledger.pdf", 2, 2, "PP", 2, 3, 7, "A2"],
    ]
    chem_rows = _sheet_rows(workbook, chem_sheet)
    assert chem_rows[0] == ["파일", "페이지", "행번호", "원소", "값"]
    assert chem_rows[1] == ["ledger.pdf", 1, 1, "C", 0.12]
    assert chem_rows[2] == ["ledger.pdf", 1, 2, "Si", None]

    review = _sheet_rows(workbook, "검토_칸")
    assert review[0] == ["파일", "페이지", "표", "행번호", "칼럼", "값", "원문", "신뢰도", "flags"]
    low_confidence = [row for row in review[1:] if row[4] == "단중"]
    assert [(row[1], row[3], row[7]) for row in low_confidence] == [(1, 2, 0.7), (2, 2, 0.7)]
    uncertain = [row for row in review[1:] if row[4] == "값"]
    assert [(row[1], row[6], row[8]) for row in uncertain] == [(1, "1.5?", "uncertain_chars"), (2, "1.5?", "uncertain_chars")]
    assert not any(row[4] == "주입중량" for row in review[1:])  # the corrected cell is not a review cell

    violations = _sheet_rows(workbook, "검증_위반")
    assert violations[0] == ["파일", "페이지", "표", "행번호", "규칙", "내용"]
    assert [(row[1], row[2], row[3], row[4]) for row in violations[1:]] == [(2, "주입 내역", 2, "row_formula")]


def test_batch_export_options_csv_zip_json_and_export_job(mock_vlm, fake_vlm, client) -> None:
    batch, _jobs = _ledger_batch(fake_vlm, client)

    response = client.get(f"/api/table-extraction-batches/{batch['id']}/export?format=xlsx&include_raw=true&include_confidence=true")
    workbook = load_workbook(io.BytesIO(response.content))
    header, first = _sheet_rows(workbook, "주입 내역")[:2]
    assert header[:9] == ["파일", "페이지", "행번호", "품명", "품명_원문", "품명_신뢰도", "단중", "단중_원문", "단중_신뢰도"]
    assert first[3:6] == ["MBS", "MBS", 0.95]
    assert header[header.index("레들") + 1] == "레들_원문" and first[header.index("레들_원문")] == "A1:1, A3:1"
    chem_header = _sheet_rows(workbook, workbook.sheetnames[1])[0]
    assert chem_header == ["파일", "페이지", "행번호", "원소", "원소_원문", "원소_신뢰도", "값", "값_신뢰도"]  # keep_raw=false

    csv_response = client.get(f"/api/table-extraction-batches/{batch['id']}/export?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"] == "application/zip"
    assert ".zip" in csv_response.headers["content-disposition"]
    archive = zipfile.ZipFile(io.BytesIO(csv_response.content))
    names = archive.namelist()
    assert names[0] == "주입 내역.csv" and names[-2:] == ["검토_칸.csv", "검증_위반.csv"] and len(names) == 4
    injection_csv = archive.read("주입 내역.csv")
    assert injection_csv.startswith(b"\xef\xbb\xbf")
    lines = injection_csv.decode("utf-8-sig").splitlines()
    assert lines[0] == "파일,페이지,행번호,품명,단중,실적,주입중량,레들"
    assert lines[1] == "ledger.pdf,1,1,MBS,2,3,6,A1;A3"

    payload = client.get(f"/api/table-extraction-batches/{batch['id']}/export?format=json").json()
    assert payload["template_name"] == "샘플 주입일지" and payload["total_count"] == 2
    assert [page["page_number"] for page in payload["pages"]] == [1, 2]
    assert [page["corrected"] for page in payload["pages"]] == [True, False]
    assert payload["pages"][0]["tables"][0]["rows"][1]["cells"]["주입중량"]["value"] == 6
    assert payload["pages"][1]["tables"][0]["rows"][1]["cells"]["주입중량"]["value"] == 7

    export_job = client.post("/api/export-jobs", json={"owner_type": "table_extraction_batch", "owner_id": batch["id"], "format": "csv"})
    assert export_job.status_code == 202, export_job.text
    completed = wait_for_export_job(client, export_job.json()["id"])
    assert completed["status"] == "completed", completed
    assert completed["filename"].endswith(".zip") and completed["content_type"] == "application/zip"
    download = client.get(f"/api/export-jobs/{completed['id']}/download")
    assert zipfile.ZipFile(io.BytesIO(download.content)).namelist() == names
    listed = client.get(f"/api/export-jobs?owner_type=table_extraction_batch&owner_id={batch['id']}").json()
    assert [job["id"] for job in listed] == [completed["id"]]
    assert client.get("/api/table-extraction-batches/missing/export").status_code == 404


def test_single_job_export_and_unfinished_job(mock_vlm, fake_vlm, client) -> None:
    _batch, jobs = _ledger_batch(fake_vlm, client)
    response = client.get(f"/api/table-extraction-jobs/{jobs[1]['job_id']}/export?format=xlsx")
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames[0] == "주입 내역" and workbook.sheetnames[-2:] == ["검토_칸", "검증_위반"]
    assert [row[1] for row in _sheet_rows(workbook, "주입 내역")[1:]] == [2, 2]

    payload = client.get(f"/api/table-extraction-jobs/{jobs[0]['job_id']}/export?format=json").json()
    assert payload["job_id"] == jobs[0]["job_id"] and payload["corrected"] is True
    assert [table["table_key"] for table in payload["tables"]] == ["injection_rows", "chem"]

    def explode(table, image_inputs, *, has_region_crop):
        raise RuntimeError("model exploded")

    fake_vlm.set(extract_table=explode)
    template = _create_template(client, name="broken", tables=[_injection_table()])
    document = _upload_image(client)
    created = client.post("/api/table-extraction-jobs", json={"document_id": document["document_id"], "template_id": template["id"]}).json()
    assert client.get(f"/api/table-extraction-jobs/{created['job_id']}/export?format=xlsx").status_code == 409
    assert client.get("/api/table-extraction-jobs/missing/export").status_code == 404


def test_sheet_names_are_sanitized_and_deduplicated() -> None:
    used = {"검토_칸"}
    assert unique_sheet_name("표", used) == "표"
    assert unique_sheet_name("표", used) == "표 (2)"
    assert unique_sheet_name("검토_칸", used) == "검토_칸 (2)"
    long_name = unique_sheet_name("가" * 40, used)
    assert long_name == "가" * 31
    assert unique_sheet_name("가" * 40, used) == "가" * 27 + " (2)"
    assert unique_sheet_name("  ", used, "fallback") == "fallback"


# --- workflow node --------------------------------------------------------------------------


def _table_workflow(template_id: str, classifier_id: str, schema_id: str, table_keys: list[str] | None = None) -> dict[str, Any]:
    table_config: dict[str, Any] = {"template_id": template_id}
    if table_keys is not None:
        table_config["table_keys"] = table_keys
    return {
        "nodes": [
            {"id": "input", "data": {"kind": "input"}},
            {"id": "classifier", "data": {"kind": "classifier", "config": {"classifier_id": classifier_id}}},
            {"id": "branch", "data": {"kind": "branch"}},
            {"id": "table_contract", "data": {"kind": "table", "label": "표 추출", "config": table_config}},
            {"id": "kie_consent", "data": {"kind": "kie", "config": {"schema_id": schema_id}}},
            {"id": "merge", "data": {"kind": "merge"}},
            {"id": "export", "data": {"kind": "export"}},
        ],
        "edges": [
            {"id": "e1", "source": "input", "target": "classifier"},
            {"id": "e2", "source": "classifier", "target": "branch"},
            {"id": "e3", "source": "branch", "target": "table_contract", "sourceHandle": "class:contract"},
            {"id": "e4", "source": "branch", "target": "kie_consent", "sourceHandle": "class:consent_form"},
            {"id": "e5", "source": "branch", "target": "merge", "sourceHandle": "unknown"},
            {"id": "e6", "source": "table_contract", "target": "merge"},
            {"id": "e7", "source": "kie_consent", "target": "merge"},
            {"id": "e8", "source": "merge", "target": "export"},
        ],
    }


def test_workflow_table_node_under_branch_runs_per_page_and_exports_table_sheet(mock_vlm, fake_vlm, client) -> None:
    def classify(classes, allow_unknown, image_paths):
        class_name = "consent_form" if "page_3" in str(image_paths[0]) else "contract"
        return {"status": "classified", "class_name": class_name, "confidence": 0.9, "reason": "seen", "evidence": []}

    def extract_table(table, image_inputs, *, has_region_crop):
        conf = 0.5 if "page_2" in image_inputs[0]["path"] else 0.95
        return {
            "headers_seen": [],
            "rows": [
                _row(1, ("MBS", conf), ("2", conf), ("3", conf), ("6", conf), ("A1", conf)),
                _row(2, ("PP", conf), ("1", conf), ("4", conf), ("4", conf), ("", conf)),
            ],
            "notes": None,
        }

    fake_vlm.set(classify_document=classify, extract_table=extract_table)
    template = _create_template(client, name="workflow ledger", tables=[_injection_table(), _chem_table()])
    classifier = create_document_classifier(client, name="table_workflow_classifier")
    schema = create_schema(client, name="table_workflow_schema")

    missing = client.post("/api/workflows", json={"name": "bad", "definition": _table_workflow("missing", classifier["id"], schema["id"])})
    assert missing.status_code == 422 and "Table node table_contract must select a saved table template" in missing.text
    unknown_key = client.post(
        "/api/workflows", json={"name": "bad", "definition": _table_workflow(template["id"], classifier["id"], schema["id"], ["nope"])}
    )
    assert unknown_key.status_code == 422 and "Unknown table key(s): nope" in unknown_key.text

    workflow = client.post(
        "/api/workflows",
        json={"name": "표 워크플로우", "definition": _table_workflow(template["id"], classifier["id"], schema["id"], ["injection_rows"])},
    )
    assert workflow.status_code == 200, workflow.text
    document = upload_pdf(client, filename="mixed.pdf", page_count=3)
    started = client.post(f"/api/workflows/{workflow.json()['id']}/runs/from-documents", json={"document_ids": [document["document_id"]]})
    assert started.status_code == 200, started.text
    run = client.get(f"/api/workflow-runs/{started.json()['id']}").json()
    items = run["items"]
    assert [item["result"]["branch_path"] for item in items] == ["class:contract", "class:contract", "class:consent_form"]
    assert [item["status"] for item in items[:2]] == ["completed", "needs_review"]

    first = items[0]["result"]["node_results"]["table_contract"]
    assert first["kind"] == "table" and first["status"] == "completed" and first["template_id"] == template["id"]
    assert first["tables"] == [{"key": "injection_rows", "name": "주입 내역", "status": "ok", "row_count": 2, "review_cells": 0, "violations": 0}]
    second = items[1]["result"]["node_results"]["table_contract"]
    assert second["status"] == "needs_review" and second["tables"][0]["review_cells"] == 10
    job = client.get(f"/api/table-extraction-jobs/{first['job_id']}").json()
    assert (job["page_number"], job["table_keys"], job["result_id"]) == (1, ["injection_rows"], first["result_id"])
    assert [table["key"] for table in job["definition_snapshot"]["tables"]] == ["injection_rows", "chem"]
    assert "table_contract" not in items[2]["result"]["node_results"]

    xlsx = client.get(f"/api/workflow-runs/{run['id']}/export?format=xlsx")
    assert xlsx.status_code == 200, xlsx.text
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Export", "주입 내역", "검토_칸", "검증_위반"]
    export_rows = _sheet_rows(workbook, "Export")
    assert "table_injection_rows_rows" in export_rows[0] and any(column.startswith("kie_") for column in export_rows[0])
    table_rows = _sheet_rows(workbook, "주입 내역")
    assert table_rows[0] == ["파일", "페이지", "행번호", "품명", "단중", "실적", "주입중량", "레들"]
    assert [row[:3] for row in table_rows[1:]] == [["mixed.pdf", 1, 1], ["mixed.pdf", 1, 2], ["mixed.pdf", 2, 1], ["mixed.pdf", 2, 2]]
    assert table_rows[1][3:] == ["MBS", 2, 3, 6, "A1"]
    assert len(_sheet_rows(workbook, "검토_칸")) == 1 + 10

    payload = client.get(f"/api/workflow-runs/{run['id']}/export?format=json").json()
    (group,) = payload["tables"]
    assert (group["template_id"], group["template_name"], group["table_key"]) == (template["id"], "workflow ledger", "injection_rows")
    assert [row["page_number"] for row in group["rows"]] == [1, 1, 2, 2]
    assert group["rows"][0]["values"] == {"품명": "MBS", "단중": 2, "실적": 3, "주입중량": 6, "레들": ["A1"]}
    assert payload["rows"][0]["table_injection_rows_rows"] == 2

    csv_text = client.get(f"/api/workflow-runs/{run['id']}/export?format=csv").text
    assert "table_injection_rows_rows" in csv_text.splitlines()[0]


def test_workflow_without_table_nodes_exports_unchanged(mock_vlm, client) -> None:
    from tests.support.api import create_required_field_checklist, workflow_definition

    schema = create_schema(client, name="plain_schema")
    classifier = create_document_classifier(client, name="plain_classifier")
    checklist = create_required_field_checklist(client, name="plain_checklist")
    workflow = client.post("/api/workflows", json={"name": "plain", "definition": workflow_definition(schema["id"], classifier["id"], checklist["id"])})
    document = upload_pdf(client, page_count=1)
    run = client.post(f"/api/workflows/{workflow.json()['id']}/runs/from-documents", json={"document_ids": [document["document_id"]]}).json()
    assert "tables" not in client.get(f"/api/workflow-runs/{run['id']}/export?format=json").json()
    workbook = load_workbook(io.BytesIO(client.get(f"/api/workflow-runs/{run['id']}/export?format=xlsx").content))
    assert workbook.sheetnames == ["Export"]


# --- recommendations ------------------------------------------------------------------------


def test_recommendation_with_mock_vlm_returns_a_three_column_draft(mock_vlm, client) -> None:
    document = _upload_image(client)
    region = {"x": 0.1, "y": 0.2, "width": 0.5, "height": 0.3}
    response = client.post(
        "/api/table-templates/recommendations",
        json={"document_id": document["document_id"], "page_number": 1, "region": region, "table_description": "주입 기록"},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    table = payload["table"]
    assert (table["key"], table["name"], table["description"]) == ("work_rows", "작업 내역", "주입 기록")
    assert [(column["key"], column["type"], column["header_path"]) for column in table["columns"]] == [
        ("품명", "text", ["품명"]),
        ("수량", "number", ["실적", "수량"]),
        ("작업일", "date", ["작업일"]),
    ]
    assert table["region"] == region and table["end_hint"] == "'합계' 줄 위까지"
    assert table["row_count"] == {"mode": "max", "value": 10}
    assert payload["reasoning"]
    # The draft is a valid TableSpec: it can be saved as is.
    assert client.post("/api/table-templates", json={"name": "from draft", "tables": [table]}).status_code == 200


def test_recommendation_crops_the_region_and_sanitizes_the_model_draft(mock_vlm, fake_vlm, client) -> None:
    seen: list[dict[str, Any]] = []

    def recommend_table(image_inputs, *, table_description, has_region_crop):
        sizes = []
        for image_input in image_inputs:
            with Image.open(io.BytesIO(read_storage_bytes(image_input["path"]))) as image:
                sizes.append(image.size)
        seen.append({"sizes": sizes, "description": table_description, "has_region_crop": has_region_crop})
        return {
            "name": "  주입   내역 ",
            "key": "주입 내역/표",
            "description": None,
            "columns": [
                {"key": "수량", "header_path": ["주입수량", "실적", "", "a", "b"], "type": "number", "description": "", "choices": [], "checkbox_options": []},
                {"key": "수량", "header_path": [], "type": "money", "description": None, "choices": [], "checkbox_options": []},
                {"key": " ", "header_path": ["잔량"], "type": "choice", "description": None, "choices": [], "checkbox_options": []},
                {"key": "판정", "header_path": [], "type": "choice", "description": None, "choices": ["완", "완", "미"], "checkbox_options": []},
                {"key": "레들", "header_path": [], "type": "checkbox", "description": None, "choices": ["x"], "checkbox_options": ["A1", "A2"]},
                "garbage",
            ],
            "end_hint": "  ",
            "row_unit_hint": "주입 1건",
            "row_count_guess": 0,
            "reasoning": "",
        }

    fake_vlm.set(recommend_table=recommend_table)
    document = _upload_image(client, width=200, height=300)
    response = client.post(
        "/api/table-templates/recommendations",
        json={"document_id": document["document_id"], "region": {"x": 0.1, "y": 0.2, "width": 0.5, "height": 0.3}},
    )
    assert response.status_code == 200, response.text
    assert seen[0]["has_region_crop"] is True and seen[0]["description"] is None
    # region + 5% padding of the page: x 0.05~0.65 → 120px, y 0.15~0.55 → 120px; plus the downscaled page.
    assert seen[0]["sizes"] == [(120, 120), (200, 300)]
    table = response.json()["table"]
    assert (table["key"], table["name"], table["end_hint"], table["row_unit_hint"]) == ("주입_내역_표", "주입 내역", None, "주입 1건")
    assert table["row_count"] == {"mode": "unknown", "value": None}
    columns = [(column["key"], column["type"], column["header_path"], column["choices"], column["checkbox_options"]) for column in table["columns"]]
    assert columns == [
        ("수량", "number", ["주입수량", "실적", "a", "b"], [], []),
        ("수량_2", "text", [], [], []),
        ("잔량", "text", ["잔량"], [], []),
        ("판정", "choice", [], ["완", "미"], []),
        ("레들", "checkbox", [], [], ["A1", "A2"]),
    ]
    assert response.json()["reasoning"] is None

    whole_page = client.post("/api/table-templates/recommendations", json={"document_id": document["document_id"], "table_description": " 주입 "})
    assert whole_page.status_code == 200
    assert seen[1] == {"sizes": [(200, 300)], "description": "주입", "has_region_crop": False}


def test_recommendation_errors(mock_vlm, fake_vlm, client) -> None:
    document = _upload_image(client)
    assert client.post("/api/table-templates/recommendations", json={"document_id": "missing"}).status_code == 404
    assert client.post("/api/table-templates/recommendations", json={"document_id": document["document_id"], "page_number": 2}).status_code == 422

    fake_vlm.set(recommend_table=lambda image_inputs, *, table_description, has_region_crop: {"name": "x", "columns": []})
    assert client.post("/api/table-templates/recommendations", json={"document_id": document["document_id"]}).status_code == 502

    def explode(image_inputs, *, table_description, has_region_crop):
        raise RuntimeError("model exploded")

    fake_vlm.set(recommend_table=explode)
    failed = client.post("/api/table-templates/recommendations", json={"document_id": document["document_id"]})
    assert failed.status_code == 400 and "model exploded" in failed.text
