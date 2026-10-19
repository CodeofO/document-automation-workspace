"""Table extraction exports (design §5 Export): long-table sheets, review cells, rule violations.

Shared by the table batch export, the single-job export (review screen) and the workflow run
export (table nodes). Everything here works on stored ``TableResult`` dicts:

* the effective tables of a result are ``corrected_tables`` when present, else ``tables``;
* one sheet/CSV per table (grouped by table key, and by template too for workflow runs):
  ``파일, 페이지, 행번호`` then one column per ColumnSpec (optionally ``<칼럼>_원문`` /
  ``<칼럼>_신뢰도`` next to it). Checkbox values are joined like ``"A1;A3"``;
* ``검토_칸``: cells with confidence < 0.8 or a review-relevant flag;
* ``검증_위반``: row violations, failed table checks, failed tables and failed/unfinished pages.
"""

import csv
import io
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook

from app.domain.table_normalization import REVIEW_CONFIDENCE_THRESHOLD
from app.domain.table_spec import TableSpec, TableTemplateDefinition, load_definition

REVIEW_SHEET_NAME = "검토_칸"
VIOLATION_SHEET_NAME = "검증_위반"
REVIEW_HEADERS = ["파일", "페이지", "표", "행번호", "칼럼", "값", "원문", "신뢰도", "flags"]
VIOLATION_HEADERS = ["파일", "페이지", "표", "행번호", "규칙", "내용"]
ROW_META_HEADERS = ["파일", "페이지", "행번호"]
RAW_SUFFIX = "_원문"
CONFIDENCE_SUFFIX = "_신뢰도"
# Flags that mark a cell's kind or a confirmed state rather than a doubt; on their own they do
# not put a cell on the review sheet (an empty *required* cell shows up as a row violation).
REVIEW_IGNORED_FLAGS = frozenset({"empty", "checkbox", "choice_matched", "corrected"})
SHEET_NAME_LIMIT = 31
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def effective_tables(result: Any) -> list[dict[str, Any]]:
    """``corrected_tables`` (human-reviewed copy) when present, else the model ``tables``."""
    if result is None:
        return []
    tables = result.corrected_tables if result.corrected_tables is not None else result.tables
    return [table for table in tables or [] if isinstance(table, dict)]


def export_value(value: Any) -> Any:
    """Cell value for a spreadsheet/CSV: checkbox lists joined with ``;``."""
    if isinstance(value, list):
        return ";".join(str(item) for item in value)
    return value


def _cell(row: dict[str, Any], column: str) -> dict[str, Any]:
    cells = row.get("cells") if isinstance(row.get("cells"), dict) else {}
    cell = cells.get(column)
    return cell if isinstance(cell, dict) else {}


def is_review_cell(cell: dict[str, Any]) -> bool:
    confidence = cell.get("confidence")
    if isinstance(confidence, int | float) and confidence < REVIEW_CONFIDENCE_THRESHOLD:
        return True
    return any(flag not in REVIEW_IGNORED_FLAGS for flag in cell.get("flags") or [])


def review_cells(table: dict[str, Any]) -> list[tuple[dict[str, Any], str, dict[str, Any]]]:
    """``(row, column label, cell)`` of every cell that belongs on the review sheet."""
    found: list[tuple[dict[str, Any], str, dict[str, Any]]] = []
    columns = [str(column) for column in table.get("columns") or []]
    for row in table.get("rows") or []:
        if not isinstance(row, dict):
            continue
        row_columns = columns or list((row.get("cells") or {}).keys())
        for column in row_columns:
            cell = _cell(row, column)
            if cell and is_review_cell(cell):
                found.append((row, column, cell))
        for index, cell in enumerate(row.get("extra_cells") or [], start=1):
            if isinstance(cell, dict):
                found.append((row, f"(추가 칸 {index})", cell))
    return found


def table_violations(table: dict[str, Any]) -> list[tuple[int | None, str, str]]:
    """``(row_index, rule, message)``: row violations, failed table checks and a failed table."""
    found: list[tuple[int | None, str, str]] = []
    if table.get("status") == "failed":
        found.append((None, "table_failed", str(table.get("error") or "Table extraction failed")))
    for row in table.get("rows") or []:
        if not isinstance(row, dict):
            continue
        for violation in row.get("violations") or []:
            if isinstance(violation, dict):
                found.append((row.get("row_index"), str(violation.get("rule") or ""), str(violation.get("message") or "")))
    for check in table.get("table_checks") or []:
        if isinstance(check, dict) and check.get("status") == "failed" and check.get("rule") == "row_count":
            # Row-level rules (row_formula/unique) are already listed per row above.
            found.append((None, "row_count", str(check.get("message") or "")))
    return found


def table_node_summary(table: dict[str, Any]) -> dict[str, Any]:
    """Per-table summary stored in workflow ``node_results`` (details stay in the result)."""
    return {
        "key": table.get("table_key"),
        "name": table.get("table_name"),
        "status": table.get("status"),
        "row_count": table.get("row_count", 0),
        "review_cells": len(review_cells(table)),
        "violations": len(table_violations(table)),
    }


# --- sheet naming ---------------------------------------------------------------------------


def sanitize_sheet_name(name: Any, fallback: str = "Sheet") -> str:
    text = _INVALID_SHEET_CHARS.sub("_", str(name or "")).strip().strip("'")
    return (text or fallback)[:SHEET_NAME_LIMIT]


def unique_sheet_name(name: Any, used: set[str], fallback: str = "Sheet") -> str:
    """Excel-safe (≤31 chars, no ``[]:*?/\\``) name not in ``used`` (case-insensitive); adds ``(2)``, ``(3)``…"""
    base = sanitize_sheet_name(name, fallback)
    taken = {item.casefold() for item in used}
    candidate = base
    counter = 2
    while candidate.casefold() in taken:
        suffix = f" ({counter})"
        candidate = base[: SHEET_NAME_LIMIT - len(suffix)] + suffix
        counter += 1
    used.add(candidate)
    return candidate


def safe_filename(name: str) -> str:
    text = _INVALID_FILENAME_CHARS.sub("_", name).strip(". ")
    return text or "table"


# --- collection -----------------------------------------------------------------------------


@dataclass
class TableSheet:
    table_key: str
    table_name: str
    template_id: str | None = None
    template_name: str | None = None
    columns: list[str] = field(default_factory=list)
    raw_columns: set[str] = field(default_factory=set)
    rows: list[dict[str, Any]] = field(default_factory=list)

    def add_columns(self, columns: Iterable[str], spec: TableSpec | None) -> None:
        keep_raw = {column.key: column.keep_raw for column in spec.columns} if spec else {}
        for column in columns:
            if column not in self.columns:
                self.columns.append(column)
            if keep_raw.get(column, True):
                self.raw_columns.add(column)


@dataclass(frozen=True)
class PageSource:
    """Where a page's tables came from (one job / batch item / workflow run item)."""

    filename: str | None
    page_number: int | None
    document_id: str | None = None
    job_id: str | None = None
    workflow_run_item_id: str | None = None


class TableExportCollector:
    """Accumulates the rows of many pages, grouped into one sheet per table (per template for workflows)."""

    def __init__(self, *, include_raw: bool = False, include_confidence: bool = False, group_by_template: bool = False) -> None:
        self.include_raw = include_raw
        self.include_confidence = include_confidence
        self.group_by_template = group_by_template
        self.sheets: dict[tuple[str | None, str], TableSheet] = {}
        self.review_rows: list[list[Any]] = []
        self.violation_rows: list[list[Any]] = []

    def __bool__(self) -> bool:
        return bool(self.sheets)

    def _group(self, template_id: str | None, table_key: str) -> tuple[str | None, str]:
        return (template_id if self.group_by_template else None, table_key)

    def declare(self, definition: TableTemplateDefinition, *, template_id: str | None = None, template_name: str | None = None) -> None:
        """Create the sheets of a definition up front (keeps definition order, empty tables still get a sheet)."""
        for spec in definition.tables:
            self._sheet(spec.key, spec.name, spec, template_id, template_name).add_columns(spec.column_keys, spec)

    def _sheet(self, table_key: str, table_name: str, spec: TableSpec | None, template_id: str | None, template_name: str | None) -> TableSheet:
        group = self._group(template_id, table_key)
        sheet = self.sheets.get(group)
        if sheet is None:
            sheet = TableSheet(table_key=table_key, table_name=table_name or table_key, template_id=template_id, template_name=template_name)
            self.sheets[group] = sheet
        return sheet

    def add_failure(self, source: PageSource, rule: str, message: str) -> None:
        self.violation_rows.append([source.filename, source.page_number, None, None, rule, message])

    def add_tables(
        self,
        source: PageSource,
        tables: list[dict[str, Any]],
        definition: TableTemplateDefinition | None = None,
        *,
        template_id: str | None = None,
        template_name: str | None = None,
    ) -> None:
        specs = {spec.key: spec for spec in (definition.tables if definition else [])}
        for table in tables:
            table_key = str(table.get("table_key") or "")
            spec = specs.get(table_key)
            table_name = str(table.get("table_name") or (spec.name if spec else table_key))
            sheet = self._sheet(table_key, table_name, spec, template_id, template_name)
            columns = [str(column) for column in table.get("columns") or (spec.column_keys if spec else [])]
            sheet.add_columns(columns, spec)
            for row in table.get("rows") or []:
                if isinstance(row, dict):
                    sheet.rows.append(self._row_values(source, row, columns))
            for row, column, cell in review_cells(table):
                self.review_rows.append(
                    [
                        source.filename,
                        source.page_number,
                        table_name,
                        row.get("row_index"),
                        column,
                        export_value(cell.get("value")),
                        cell.get("raw"),
                        cell.get("confidence"),
                        ";".join(str(flag) for flag in cell.get("flags") or []),
                    ]
                )
            for row_index, rule, message in table_violations(table):
                self.violation_rows.append([source.filename, source.page_number, table_name, row_index, rule, message])

    def _row_values(self, source: PageSource, row: dict[str, Any], columns: list[str]) -> dict[str, Any]:
        values: dict[str, Any] = {
            "__source__": source,
            "__row__": row,
            ("meta", "파일"): source.filename,
            ("meta", "페이지"): source.page_number,
            ("meta", "행번호"): row.get("row_index"),
        }
        for column in columns:
            cell = _cell(row, column)
            values[("value", column)] = export_value(cell.get("value"))
            values[("raw", column)] = cell.get("raw")
            values[("confidence", column)] = cell.get("confidence")
        return values

    # --- layout ---------------------------------------------------------------------------

    def _layout(self, sheet: TableSheet) -> list[tuple[tuple[str, str], str]]:
        layout: list[tuple[tuple[str, str], str]] = [(("meta", header), header) for header in ROW_META_HEADERS]
        for column in sheet.columns:
            layout.append((("value", column), column))
            if self.include_raw and column in sheet.raw_columns:
                layout.append((("raw", column), f"{column}{RAW_SUFFIX}"))
            if self.include_confidence:
                layout.append((("confidence", column), f"{column}{CONFIDENCE_SUFFIX}"))
        return layout

    def sheet_table(self, sheet: TableSheet) -> tuple[list[str], list[list[Any]]]:
        layout = self._layout(sheet)
        return [label for _key, label in layout], [[row.get(key) for key, _label in layout] for row in sheet.rows]

    def named_sheets(self, reserved: Iterable[str] = ()) -> list[tuple[str, TableSheet]]:
        used = set(reserved) | {REVIEW_SHEET_NAME, VIOLATION_SHEET_NAME}
        return [(unique_sheet_name(sheet.table_name, used, sheet.table_key), sheet) for sheet in self.sheets.values()]

    def write_sheets(self, workbook: Workbook, *, reserved: Iterable[str] = (), replace_active: bool = False) -> None:
        """Append one sheet per table, then ``검토_칸`` and ``검증_위반``."""
        first = replace_active
        for name, sheet in self.named_sheets(reserved):
            worksheet = workbook.active if first else workbook.create_sheet()
            first = False
            worksheet.title = name
            headers, rows = self.sheet_table(sheet)
            worksheet.append(headers)
            for row in rows:
                worksheet.append([_xlsx_value(value) for value in row])
        for title, headers, rows in (
            (REVIEW_SHEET_NAME, REVIEW_HEADERS, self.review_rows),
            (VIOLATION_SHEET_NAME, VIOLATION_HEADERS, self.violation_rows),
        ):
            worksheet = workbook.active if first else workbook.create_sheet()
            first = False
            worksheet.title = title
            worksheet.append(headers)
            for row in rows:
                worksheet.append([_xlsx_value(value) for value in row])

    def xlsx_bytes(self) -> bytes:
        workbook = Workbook()
        self.write_sheets(workbook, replace_active=True)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def csv_zip_bytes(self) -> bytes:
        """One UTF-8-with-BOM CSV per table, plus ``검토_칸.csv`` and ``검증_위반.csv``."""
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            used: set[str] = set()
            for name, sheet in self.named_sheets():
                headers, rows = self.sheet_table(sheet)
                archive.writestr(_unique_csv_name(name, used), _csv_text(headers, rows))
            archive.writestr(f"{REVIEW_SHEET_NAME}.csv", _csv_text(REVIEW_HEADERS, self.review_rows))
            archive.writestr(f"{VIOLATION_SHEET_NAME}.csv", _csv_text(VIOLATION_HEADERS, self.violation_rows))
        return output.getvalue()

    def json_groups(self) -> list[dict[str, Any]]:
        """Rows grouped per table (workflow JSON export): ``values`` keyed by column."""
        groups: list[dict[str, Any]] = []
        for sheet in self.sheets.values():
            rows = []
            for values in sheet.rows:
                source: PageSource = values["__source__"]
                row: dict[str, Any] = values["__row__"]
                rows.append(
                    {
                        "filename": source.filename,
                        "page_number": source.page_number,
                        "document_id": source.document_id,
                        "job_id": source.job_id,
                        "workflow_run_item_id": source.workflow_run_item_id,
                        "row_id": row.get("row_id"),
                        "row_index": row.get("row_index"),
                        "status": row.get("status"),
                        "values": {column: _cell(row, column).get("value") for column in sheet.columns},
                        "cells": row.get("cells") or {},
                        "violations": row.get("violations") or [],
                    }
                )
            groups.append(
                {
                    "template_id": sheet.template_id,
                    "template_name": sheet.template_name,
                    "table_key": sheet.table_key,
                    "table_name": sheet.table_name,
                    "columns": list(sheet.columns),
                    "rows": rows,
                }
            )
        return groups


def _xlsx_value(value: Any) -> Any:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    return str(value)


def _csv_text(headers: list[str], rows: list[list[Any]]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return ("﻿" + output.getvalue()).encode("utf-8")


def _unique_csv_name(name: str, used: set[str]) -> str:
    base = safe_filename(name)
    candidate = f"{base}.csv"
    counter = 2
    while candidate.casefold() in used:
        candidate = f"{base} ({counter}).csv"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def job_definition(job: Any) -> TableTemplateDefinition | None:
    try:
        return load_definition(job.definition_snapshot) if job is not None else None
    except ValueError:
        return None
